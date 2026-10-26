"""
Descarga los indices MSCI de paises (y agregados) desde la API publica de MSCI
("End of Day Index Data") y los guarda en Excel.

Salida: Database/MSCI/MSCI_Indexes.xlsx

Hojas de datos (misma estructura que Yahoo_Prices.xlsx):
  Row 1 : Description
  Row 2 : Source
  Row 3 : Unit
  Row 4 : Frequency
  Row 5 : SA
  Row 6 : ISO3 / Code
  Row 7+: Datos diarios (col A = fecha)

  Price_USD  - indice de precio (STRD) en USD
  Price_LOC  - indice de precio (STRD) en moneda local
  NetTR_USD  - indice de retorno neto (NETR) en USD
  Index_List - nombre, codigo MSCI, ISO3, clasificacion MSCI, primera / ultima fecha

La lista de paises y su clasificacion (Developed / Emerging / Frontier & Standalone)
se lee en cada corrida de las listas oficiales de MSCI, asi que reclasificaciones
quedan reflejadas automaticamente.

Uso:
    python MSCI_Download.py                 # incremental (o completo si no existe el Excel)
    python MSCI_Download.py --full          # rehace todo desde --start
    python MSCI_Download.py --start 2010-01-01 --full

Incremental: lee el Excel existente y baja solo desde la ultima fecha (con OVERLAP_DAYS
de solapamiento, para tomar revisiones). Los indices nuevos en el universo se bajan con
historia completa; los que MSCI saca de sus listas se conservan con su historia.
Corre todos los dias via Task Scheduler (MSCI_scheduled.bat, log en Scripts/logs/).
"""

import os
import sys
import time
import argparse
import warnings
from datetime import date
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

# ── Salida ────────────────────────────────────────────────────────────────────

DATABASE = Path(__file__).parent.parent
OUTPUT = DATABASE / "MSCI" / "MSCI_Indexes.xlsx"
ISO_TABLE = DATABASE / "ISO_Master_Table.xlsx"
OVERLAP_DAYS = 10

# ── API MSCI ──────────────────────────────────────────────────────────────────

BASE = "https://app2.msci.com/products/service/index/indexmaster"
HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/128 Safari/537.36"}

# index_market de la API -> clasificacion MSCI
MARKETS = {16384: "Developed", 8192: "Emerging", 4096: "Frontier / Standalone"}

# Variantes que no son el indice estandar del pais
EXCLUDE = ("DOMESTIC", "ADR", "GROWTH", "VALUE", "CUSTOM", "DECREMENT", "REBASED", "INCLUSION", "RMB",
           "PARTIAL", "CONNECT", "CHINA A", "CHINA 50", "NIHONKABU", "MSCI ")

# Agregados (nombre en la API -> nombre en el Excel)
AGGREGATES = {
    "WORLD": "World", "ACWI": "ACWI", "EM (EMERGING MARKETS)": "EM", "EM ASIA": "EM Asia",
    "EM LATIN AMERICA": "EM Latin America", "EM EMEA": "EM EMEA", "EM ASEAN": "EM ASEAN",
    "FM (FRONTIER MARKETS)": "Frontier Markets", "EAFE": "EAFE", "EUROPE": "Europe", "NORTH AMERICA": "North America",
}

SHEETS = {
    "Price_USD": {"variant": "STRD", "currency": "USD", "description": "MSCI Price Index (USD)", "unit": "Index, USD"},
    "Price_LOC": {"variant": "STRD", "currency": "LOC", "description": "MSCI Price Index (local currency)", "unit": "Index, LC"},
    "NetTR_USD": {"variant": "NETR", "currency": "USD", "description": "MSCI Net Total Return Index (USD)", "unit": "Index, USD"},
}

# Nombres MSCI que no coinciden con ISO_Master_Table
ISO_FIX = {
    "KOREA": "KOR", "CZECH REPUBLIC": "CZE", "IVORY COAST": "CIV", "UNITED KINGDOM": "GBR", "USA": "USA",
    "RUSSIA": "RUS", "VIETNAM": "VNM", "TAIWAN": "TWN", "HONG KONG": "HKG", "LATIVA": "LVA", "TURKEY": "TUR",
    "BOSNIA AND HERZEGOVINA": "BIH", "TRINIDAD AND TOBAGO": "TTO", "UNITED ARAB EMIRATES": "ARE", "EGYPT": "EGY",
    "IRAN": "IRN", "SLOVAKIA": "SVK",
}
NAME_FIX = {"LATIVA": "Latvia", "USA": "USA", "UNITED KINGDOM": "UK", "UNITED ARAB EMIRATES": "UAE"}


def get(url, params, tries=4):
    for k in range(tries):
        try:
            r = requests.get(url, params=params, headers=HEADERS, timeout=90)
            r.raise_for_status()
            return r.json()
        except Exception:
            if k == tries - 1:
                raise
            time.sleep(2 * (k + 1))


def index_universe():
    """Country indexes per MSCI classification plus selected aggregates."""
    rows = []
    for market, cls in MARKETS.items():
        js = get(f"{BASE}/searchIndexes", {"index_market": market, "index_scope": "Country", "index_size": 12,
                                           "index_style": "None", "index_suite": "C"})
        for x in js["indexes"]:
            name = x["index_name"]
            if any(e in name.upper() for e in EXCLUDE):
                continue
            rows.append({"msci_name": name, "code": int(x["msci_index_code"]), "classification": cls, "type": "Country"})
        js = get(f"{BASE}/searchIndexes", {"index_market": market, "index_scope": "Region", "index_size": 12,
                                           "index_style": "None", "index_suite": "C"})
        for x in js["indexes"]:
            if x["index_name"] in AGGREGATES:
                rows.append({"msci_name": x["index_name"], "code": int(x["msci_index_code"]), "classification": cls,
                             "type": "Aggregate"})
    js = get(f"{BASE}/searchIndexes", {"index_market": 24576, "index_scope": "Region", "index_size": 12,
                                       "index_style": "None", "index_suite": "C"})
    for x in js["indexes"]:
        if x["index_name"] in AGGREGATES:
            rows.append({"msci_name": x["index_name"], "code": int(x["msci_index_code"]), "classification": "All Country",
                         "type": "Aggregate"})
    df = pd.DataFrame(rows).drop_duplicates("code")

    iso = pd.read_excel(ISO_TABLE)
    iso_map = {str(c).upper(): i for c, i in zip(iso["Country"], iso["ISO3"])}

    def iso3(r):
        if r["type"] == "Aggregate":
            return AGGREGATES[r["msci_name"]].upper().replace(" ", "_")
        return ISO_FIX.get(r["msci_name"], iso_map.get(r["msci_name"], ""))

    def nice(r):
        if r["type"] == "Aggregate":
            return AGGREGATES[r["msci_name"]]
        return NAME_FIX.get(r["msci_name"], r["msci_name"].title())

    df["ISO3"] = df.apply(iso3, axis=1)
    df["Name"] = df.apply(nice, axis=1)
    order = {"Aggregate": 0, "Country": 1}
    cls_order = {"All Country": 0, "Developed": 1, "Emerging": 2, "Frontier / Standalone": 3}
    df = df.sort_values(by=["type", "classification", "Name"],
                        key=lambda s: s.map(order) if s.name == "type" else (s.map(cls_order) if s.name == "classification" else s))
    return df.reset_index(drop=True)


def levels(code, variant, currency, start, end):
    js = get(f"{BASE}/getLevelDataForGraph", {"currency_symbol": currency, "index_variant": variant,
                                              "start_date": start.strftime("%Y%m%d"), "end_date": end.strftime("%Y%m%d"),
                                              "data_frequency": "DAILY", "index_codes": code})
    pts = js.get("indexes", {}).get("INDEX_LEVELS", [])
    if not pts:
        return pd.Series(dtype=float)
    s = pd.Series({pd.Timestamp(str(p["calc_date"])): p["level_eod"] for p in pts}, name=code)
    return s.sort_index()


def download_sheet(codes, cfg, starts, end):
    """starts: code -> fecha de inicio. Devuelve (DataFrame, codigos con error)."""
    failed = []

    def one(code):
        try:
            return code, levels(code, cfg["variant"], cfg["currency"], starts[code], end)
        except Exception as e:
            print(f"   [!] {code}: {e}")
            failed.append(code)
            return code, pd.Series(dtype=float)

    with ThreadPoolExecutor(max_workers=6) as ex:
        res = dict(ex.map(one, codes))
    df = pd.DataFrame({c: res[c] for c in codes}, columns=list(codes))
    return df.sort_index(), failed


def read_existing():
    """Lee el Excel actual: (universo segun Index_List, {hoja: DataFrame con columnas = codigo MSCI}) o None."""
    if not OUTPUT.exists():
        return None
    try:
        xl = pd.ExcelFile(OUTPUT)
        lst = xl.parse("Index_List")
        old_u = pd.DataFrame({"msci_name": lst["MSCI name"], "code": lst["MSCI code"].astype(int),
                              "classification": lst["Classification"], "type": lst["Type"],
                              "ISO3": lst["ISO3"].fillna(""), "Name": lst["Name"]})
        # La fila 6 de las hojas tiene ISO3 (o el nombre si no hay ISO3)
        label_to_code = {(r.ISO3 or r.Name): r.code for r in old_u.itertuples()}
        data = {}
        for sh in SHEETS:
            if sh not in xl.sheet_names:
                continue
            df = xl.parse(sh, header=5, index_col=0)
            df.index = pd.to_datetime(df.index, errors="coerce")
            df = df[df.index.notna()]
            df = df[[c for c in df.columns if c in label_to_code]].rename(columns=label_to_code)
            data[sh] = df.apply(pd.to_numeric, errors="coerce").sort_index()
        return old_u, data
    except Exception as e:
        print(f"   [!] No se pudo leer {OUTPUT.name} ({e}); se hace descarga completa")
        return None


# ── Excel ─────────────────────────────────────────────────────────────────────

NAVY, GREY_L = "0E2841", "E8E8E8"
HEAD_FONT = Font(name="Raleway", bold=True, color="FFFFFF", size=9)
META_FONT = Font(name="Raleway", size=9, color="414042")
BODY_FONT = Font(name="Raleway", size=9)


def write_data_sheet(wb, name, df, universe, cfg):
    ws = wb.create_sheet(name)
    u = universe.set_index("code")
    meta = [
        ("Description", [f"{cfg['description']} - {u.loc[c, 'Name']} ({u.loc[c, 'classification']})" for c in df.columns]),
        ("Source", ["MSCI End of Day Index Data"] * df.shape[1]),
        ("Unit", [cfg["unit"]] * df.shape[1]),
        ("Frequency", ["Daily"] * df.shape[1]),
        ("SA", ["NSA"] * df.shape[1]),
        ("ISO3", [u.loc[c, "ISO3"] or u.loc[c, "Name"] for c in df.columns]),
    ]
    for r, (label, vals) in enumerate(meta, start=1):
        ws.cell(r, 1, label).font = HEAD_FONT if r == 6 else META_FONT
        for j, v in enumerate(vals, start=2):
            c = ws.cell(r, j, v)
            c.font = HEAD_FONT if r == 6 else META_FONT
            if r == 6:
                c.fill = PatternFill("solid", fgColor=NAVY)
                c.alignment = Alignment(horizontal="center")
    ws.cell(6, 1).fill = PatternFill("solid", fgColor=NAVY)
    for i, (d, row) in enumerate(df.iterrows(), start=7):
        ws.cell(i, 1, d.to_pydatetime()).number_format = "yyyy-mm-dd"
        for j, v in enumerate(row.values, start=2):
            if pd.notna(v):
                ws.cell(i, j, float(v))
    ws.column_dimensions["A"].width = 12
    for j in range(2, df.shape[1] + 2):
        ws.column_dimensions[get_column_letter(j)].width = 11
    ws.freeze_panes = "B7"


def write_index_list(wb, universe, data):
    ws = wb.create_sheet("Index_List", 0)
    cols = ["Name", "ISO3", "MSCI code", "MSCI name", "Classification", "Type", "First date (USD)", "Last date (USD)"]
    ws.append(cols)
    px = data.get("Price_USD")
    for r in universe.itertuples():
        first = last = None
        if px is not None and r.code in px:
            s = px[r.code].dropna()
            if len(s):
                first, last = s.index[0].to_pydatetime(), s.index[-1].to_pydatetime()
        ws.append([r.Name, r.ISO3, r.code, r.msci_name, r.classification, r.type, first, last])
    for c in ws[1]:
        c.font = HEAD_FONT; c.fill = PatternFill("solid", fgColor=NAVY)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
            if c.row % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GREY_L)
        row[6].number_format = row[7].number_format = "yyyy-mm-dd"
    for col, w in zip("ABCDEFGH", (22, 18, 11, 30, 22, 11, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--start", default="1999-01-01")
    ap.add_argument("--full", action="store_true", help="ignora el Excel existente y baja todo desde --start")
    args = ap.parse_args()
    start, end = pd.Timestamp(args.start), pd.Timestamp(date.today())

    print(f"MSCI Download - {date.today():%Y-%m-%d}")
    existing = None if args.full else read_existing()
    old_u, old_data = existing if existing else (None, {})
    print("Modo:", "incremental" if existing else "completo")

    print("Leyendo universo de indices MSCI...")
    universe = index_universe()
    for cls, g in universe.groupby("classification", sort=False):
        print(f"   {cls}: {len(g)}")
    if (universe["type"].eq("Country") & universe["ISO3"].eq("")).any():
        print("   [!] sin ISO3:", universe.loc[universe["ISO3"].eq(""), "msci_name"].tolist())

    live = list(universe["code"])
    if old_u is not None:
        added = universe.loc[~universe["code"].isin(old_u["code"]), "Name"].tolist()
        if added:
            print("   Nuevos en el universo:", added)
        # Indices que MSCI saco de sus listas: se conservan con su historia, sin actualizar
        gone = old_u[~old_u["code"].isin(universe["code"])]
        if len(gone):
            print("   Fuera del universo actual (se conservan):", gone["Name"].tolist())
            universe = pd.concat([universe, gone], ignore_index=True)

    data, n_failed, changed = {}, 0, existing is None
    for sh in SHEETS:
        cfg = SHEETS[sh]
        old = old_data.get(sh, pd.DataFrame())
        starts = {}
        for c in live:
            s = old[c].dropna() if c in old else pd.Series(dtype=float)
            starts[c] = max(start, s.index[-1] - pd.Timedelta(days=OVERLAP_DAYS)) if len(s) else start
        n_full = sum(starts[c] == start for c in live)
        t0 = time.time()
        print(f"Descargando {sh} ({cfg['variant']}, {cfg['currency']}): {len(live)} indices"
              + (f", {n_full} con historia completa" if existing and n_full else "") + "...")
        new, failed = download_sheet(live, cfg, starts, end)
        n_failed += len(failed)
        merged = new.combine_first(old) if len(old) else new
        merged = merged.reindex(columns=list(universe["code"])).sort_index()
        if not changed:
            prev = old.reindex(index=merged.index, columns=merged.columns)
            changed = not merged.round(6).equals(prev.round(6))
        data[sh] = merged
        empty = [universe.set_index("code").loc[c, "Name"] for c in merged.columns if merged[c].dropna().empty]
        print(f"   {merged.shape[1]} indices, {merged.index.min():%Y-%m-%d} a {merged.index.max():%Y-%m-%d} "
              f"({time.time() - t0:.0f}s)" + (f"; sin datos: {empty}" if empty else ""))

    if n_failed == len(live) * len(SHEETS):
        print("[!] Fallaron todas las descargas; no se toca el Excel")
        sys.exit(1)
    if old_u is not None and set(universe["code"]) != set(old_u["code"]):
        changed = True
    if not changed:
        print("Sin datos nuevos; no se reescribe el Excel")
        return

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    write_index_list(wb, universe, data)
    for sh, df in data.items():
        write_data_sheet(wb, sh, df, universe, SHEETS[sh])
    # Guardado atomico: si el Excel esta abierto no queda un archivo a medio escribir
    tmp = OUTPUT.with_name(OUTPUT.stem + "_tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, OUTPUT)
    except PermissionError:
        print(f"[!] {OUTPUT.name} esta abierto o bloqueado; quedo guardado como {tmp.name}")
        sys.exit(1)
    print(f"Guardado: {OUTPUT}")
    if n_failed:
        print(f"[!] {n_failed} descargas con error (se mantuvo el dato previo)")
        sys.exit(2)


if __name__ == "__main__":
    main()
